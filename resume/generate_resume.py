#!/usr/bin/env python3
"""Generate Sumit Kaktwan resume as DOCX and PDF."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

OUT = Path(__file__).resolve().parent


def set_run_font(run, name="Calibri", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_bottom_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def tight_paragraph(doc, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    return p


def heading_para(doc, text):
    p = tight_paragraph(doc, space_before=8, space_after=3)
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=(31, 78, 121))
    add_bottom_border(p)
    return p


def bullet(doc, text):
    p = tight_paragraph(doc, space_before=1, space_after=1)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    run = p.add_run("•  " + text)
    set_run_font(run, size=10)
    return p


def job_header(doc, left, right):
    p = tight_paragraph(doc, space_before=6, space_after=0)
    r1 = p.add_run(left)
    set_run_font(r1, size=10.5, bold=True)
    r2 = p.add_run("\t" + right)
    set_run_font(r2, size=10)
    tab = p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
    return p


def job_sub(doc, left, right):
    p = tight_paragraph(doc, space_before=0, space_after=2)
    r1 = p.add_run(left)
    set_run_font(r1, size=10, bold=False, color=(55, 65, 81))
    r2 = p.add_run("\t" + right)
    set_run_font(r2, size=10, color=(55, 65, 81))
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
    return p


def build_docx(path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    name = tight_paragraph(doc, 0, 0)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.add_run("SUMIT KAKTWAN")
    set_run_font(r, size=18, bold=True, color=(31, 78, 121))

    contact = tight_paragraph(doc, 2, 2)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = contact.add_run(
        "Bengaluru, India  |  +91 99977 76729  |  Sumitcskaktwan@gmail.com\n"
        "linkedin.com/in/sumit-kaktwan-270633157  |  github.com/Sumit211995"
    )
    set_run_font(cr, size=10)

    heading_para(doc, "Professional Summary")
    s = tight_paragraph(doc, 2, 2)
    sr = s.add_run(
        "Frontend Engineer with 3 years of experience shipping production React.js and Next.js products—"
        "multi-role admin dashboards, warranty/claim workflows, and CMS-backed websites. Strong in TypeScript, "
        "RBAC, REST API integration, and deployments on Railway and Linux VMs. Independently owned end-to-end "
        "UI for solar manufacturing, travel, and enterprise meeting tools; mentored small frontend teams on "
        "selected projects. Targeting mid-level Frontend / Next.js roles in product-led companies."
    )
    set_run_font(sr, size=10.5)

    heading_para(doc, "Technical Skills")
    skills = [
        (
            "Languages: ",
            "JavaScript (ES6+), TypeScript, HTML5, CSS3, SCSS",
        ),
        (
            "Frontend: ",
            "React.js, Next.js, Redux, Zustand, React Hook Form, Zod, Tailwind CSS",
        ),
        (
            "Data & integration: ",
            "REST APIs, Payload CMS, PostgreSQL, Azure Functions, Microsoft Graph API (via backend)",
        ),
        (
            "Tooling & deploy: ",
            "Git, GitHub, Firebase, Railway, Linux VM, Agile / code review",
        ),
    ]
    for label, rest in skills:
        p = tight_paragraph(doc, 1, 1)
        a = p.add_run(label)
        set_run_font(a, size=10.5, bold=True)
        b = p.add_run(rest)
        set_run_font(b, size=10.5)

    heading_para(doc, "Experience")

    job_header(doc, "Emmvee Technology Pvt. Ltd. (Emmvee Photovoltaic)", "Bengaluru, India")
    job_sub(doc, "Frontend Engineer", "Sep 2025 – Present")
    for t in [
        "Owned frontend delivery for five production applications (React.js / Next.js / TypeScript): solar shop-floor (Prism), warranty & claims, Board Meeting & Committee (BMC), Roamiyo travel admin, and the Emmvee Foundation website.",
        "Prism (solar production): role-based login and dashboards for QA, Production, PPC, Admin, Super Admin, and Manager—line and module reports tailored to each role.",
        "Warranty portal (React, Zustand auth): 15+ roles (homeowner, distributor, CSR, QA, Finance, Sales, and more). Invoice/PSN transfer, claim intake, QA assignment, refund/replacement/rework tickets, multi-level approvals, and certificate download after login.",
        "BMC (Next.js, Tailwind, SCSS): full UI for Teams-connected meetings—create/approve/cancel/reject, agendas, attendance, action items, AI-generated minutes, documents, and user management. Backend consumed Microsoft Graph API.",
        "Roamiyo admin (Next.js): customers, flights, stays, transactions, support tickets, itineraries, and dashboard charts. Contributed to the RoamWithRoamiyo web app (Redux, Azure Functions, Next.js) in production.",
        "Emmvee Foundation: Next.js + Payload CMS + PostgreSQL, including customized Payload admin CSS. Deployed Foundation and BMC on Railway; Prism and Warranty on Linux VM.",
        "Delivered most UIs as sole frontend owner; guided 2–3 frontend developers on selected projects.",
    ]:
        bullet(doc, t)

    job_header(doc, "Trifler India (Aquonics Tech Service Pvt. Ltd.)", "Bengaluru, India")
    job_sub(doc, "Frontend Developer", "Oct 2023 – Aug 2025")
    for t in [
        "Built a multi-step partner onboarding dashboard with React Hook Form, Zod validation, token-based auth, Google Maps location, and Firebase event logging.",
        "Developed a Next.js and TypeScript admin dashboard for user/partner management and a blog CMS.",
        "Partnered with design and backend on features, code reviews, debugging, and maintainable UI.",
    ]:
        bullet(doc, t)

    job_header(doc, "Satyam Auto Component", "Gurugram, India")
    job_sub(doc, "Engineer — industrial engineering / operations", "May 2019 – Jul 2020")
    bullet(
        doc,
        "Productivity studies (MOST), line balancing, kaizen, and ergonomics—domain background now used on solar production (Prism) UIs.",
    )

    job_header(doc, "Tata Motors", "Pantnagar, India")
    job_sub(doc, "Junior Associate — lean manufacturing", "Aug 2014 – Apr 2019")
    bullet(
        doc,
        "Work measurement, kaizen, ergonomics, and cost reduction on the shop floor.",
    )

    heading_para(doc, "Education")
    job_header(doc, "DIT University, Dehradun", "Uttarakhand, India")
    job_sub(
        doc,
        "B.Tech, Computer Science & Engineering  |  CGPA 7.7/10",
        "Aug 2020 – Jun 2023",
    )

    heading_para(doc, "Certifications")
    p = tight_paragraph(doc, 2, 2)
    r = p.add_run(
        "JavaScript and React.js — NamasteDev  ·  Web Development — Udemy  ·  Core Java — Smart Programming"
    )
    set_run_font(r, size=10.5)

    doc.save(path)


FONT_DIR = "/usr/share/fonts/truetype/dejavu"


class ResumePDF(FPDF):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.add_font("DejaVu", "", f"{FONT_DIR}/DejaVuSans.ttf")
        self.add_font("DejaVu", "B", f"{FONT_DIR}/DejaVuSans-Bold.ttf")

    def footer(self):
        self.set_y(-12)
        self.set_font("DejaVu", "", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 8, "Sumit Kaktwan  |  Frontend Engineer  |  Page " + str(self.page_no()), align="C")


NAVY = (31, 78, 121)
BODY = (30, 30, 30)
MUTED = (70, 70, 70)


def pdf_section(pdf, title):
    pdf.ln(2)
    pdf.set_font("DejaVu", "B", 10.5)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 5, title.upper(), new_x="LMARGIN", new_y="NEXT")
    y = pdf.get_y()
    pdf.set_draw_color(*NAVY)
    pdf.set_line_width(0.4)
    pdf.line(10, y, 200, y)
    pdf.ln(1.5)
    pdf.set_text_color(*BODY)


def pdf_job(pdf, company, loc, title, dates):
    pdf.set_font("DejaVu", "B", 10)
    pdf.set_text_color(*BODY)
    pdf.cell(130, 4.5, company, align="L")
    pdf.set_font("DejaVu", "", 9.5)
    pdf.cell(0, 4.5, loc, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(*MUTED)
    pdf.cell(120, 4.5, title, align="L")
    pdf.cell(0, 4.5, dates, align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(*BODY)


def pdf_bullet(pdf, text):
    pdf.set_font("DejaVu", "", 9.5)
    pdf.set_x(12)
    pdf.multi_cell(186, 4.1, "•  " + text)
    pdf.ln(0.25)


def build_pdf(path):
    pdf = ResumePDF(format="Letter", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_margins(10, 8, 10)

    pdf.set_font("DejaVu", "B", 18)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 8, "SUMIT KAKTWAN", align="C", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("DejaVu", "", 9.5)
    pdf.set_text_color(*BODY)
    pdf.cell(0, 5, "Bengaluru, India  |  +91 99977 76729  |  Sumitcskaktwan@gmail.com", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 5, "linkedin.com/in/sumit-kaktwan-270633157  |  github.com/Sumit211995", align="C", new_x="LMARGIN", new_y="NEXT")

    pdf_section(pdf, "Professional Summary")
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(
        0,
        4.2,
        "Frontend Engineer with 3 years of experience shipping production React.js and Next.js products—"
        "multi-role admin dashboards, warranty/claim workflows, and CMS-backed websites. Strong in TypeScript, "
        "RBAC, REST API integration, and deployments on Railway and Linux VMs. Independently owned end-to-end "
        "UI for solar manufacturing, travel, and enterprise meeting tools; mentored small frontend teams on "
        "selected projects. Targeting mid-level Frontend / Next.js roles in product-led companies.",
    )

    pdf_section(pdf, "Technical Skills")
    rows = [
        ("Languages: ", "JavaScript (ES6+), TypeScript, HTML5, CSS3, SCSS"),
        ("Frontend: ", "React.js, Next.js, Redux, Zustand, React Hook Form, Zod, Tailwind CSS"),
        ("Data & integration: ", "REST APIs, Payload CMS, PostgreSQL, Azure Functions, Microsoft Graph API (via backend)"),
        ("Tooling & deploy: ", "Git, GitHub, Firebase, Railway, Linux VM, Agile / code review"),
    ]
    for label, rest in rows:
        pdf.set_font("DejaVu", "B", 9.5)
        pdf.write(4.2, label)
        pdf.set_font("DejaVu", "", 9.5)
        pdf.write(4.2, rest)
        pdf.ln(4.6)

    pdf_section(pdf, "Experience")
    pdf_job(
        pdf,
        "Emmvee Technology Pvt. Ltd. (Emmvee Photovoltaic)",
        "Bengaluru, India",
        "Frontend Engineer",
        "Sep 2025 – Present",
    )
    for t in [
        "Owned frontend delivery for five production applications (React.js / Next.js / TypeScript): solar shop-floor (Prism), warranty & claims, Board Meeting & Committee (BMC), Roamiyo travel admin, and the Emmvee Foundation website.",
        "Prism (solar production): role-based login and dashboards for QA, Production, PPC, Admin, Super Admin, and Manager—line and module reports tailored to each role.",
        "Warranty portal (React, Zustand auth): 15+ roles (homeowner, distributor, CSR, QA, Finance, Sales, and more). Invoice/PSN transfer, claim intake, QA assignment, refund/replacement/rework tickets, multi-level approvals, and certificate download after login.",
        "BMC (Next.js, Tailwind, SCSS): full UI for Teams-connected meetings—create/approve/cancel/reject, agendas, attendance, action items, AI-generated minutes, documents, and user management. Backend consumed Microsoft Graph API.",
        "Roamiyo admin (Next.js): customers, flights, stays, transactions, support tickets, itineraries, and dashboard charts. Contributed to the RoamWithRoamiyo web app (Redux, Azure Functions, Next.js) in production.",
        "Emmvee Foundation: Next.js + Payload CMS + PostgreSQL, including customized Payload admin CSS. Deployed Foundation and BMC on Railway; Prism and Warranty on Linux VM.",
        "Delivered most UIs as sole frontend owner; guided 2–3 frontend developers on selected projects.",
    ]:
        pdf_bullet(pdf, t)

    pdf_job(
        pdf,
        "Trifler India (Aquonics Tech Service Pvt. Ltd.)",
        "Bengaluru, India",
        "Frontend Developer",
        "Oct 2023 – Aug 2025",
    )
    for t in [
        "Built a multi-step partner onboarding dashboard with React Hook Form, Zod validation, token-based auth, Google Maps location, and Firebase event logging.",
        "Developed a Next.js and TypeScript admin dashboard for user/partner management and a blog CMS.",
        "Partnered with design and backend on features, code reviews, debugging, and maintainable UI.",
    ]:
        pdf_bullet(pdf, t)

    pdf_job(
        pdf,
        "Satyam Auto Component",
        "Gurugram, India",
        "Engineer — industrial engineering / operations",
        "May 2019 – Jul 2020",
    )
    pdf_bullet(
        pdf,
        "Productivity studies (MOST), line balancing, kaizen, and ergonomics—domain background now used on solar production (Prism) UIs.",
    )

    pdf_job(
        pdf,
        "Tata Motors",
        "Pantnagar, India",
        "Junior Associate — lean manufacturing",
        "Aug 2014 – Apr 2019",
    )
    pdf_bullet(pdf, "Work measurement, kaizen, ergonomics, and cost reduction on the shop floor.")

    pdf_section(pdf, "Education")
    pdf_job(
        pdf,
        "DIT University, Dehradun",
        "Uttarakhand, India",
        "B.Tech, Computer Science & Engineering  |  CGPA 7.7/10",
        "Aug 2020 – Jun 2023",
    )

    pdf_section(pdf, "Certifications")
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(
        0,
        4.5,
        "JavaScript and React.js — NamasteDev  ·  Web Development — Udemy  ·  Core Java — Smart Programming",
    )

    pdf.output(path)


if __name__ == "__main__":
    docx_path = OUT / "Sumit_Kaktwan_Frontend_Resume.docx"
    pdf_path = OUT / "Sumit_Kaktwan_Frontend_Resume.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(docx_path)
    print(pdf_path)
