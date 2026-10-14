#!/usr/bin/env python3
"""Generate Tata Motors internship research report (DOCX + PDF)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent
IMG = ROOT / "images"
OUT = ROOT / "output"

STUDENT = {
    "name": "Nachiket L Karoshi",
    "department": "Diploma in Electronics and Communication Engineering",
    "institution": "KLE Society's C B Kolli Polytechnic, Haveri",
    "internship_org": "Tata Motors Limited",
    "academic_year": "2026–2027",
    "submission_date": date.today().strftime("%d %B %Y"),
}

NAVY = (0, 51, 102)
ACCENT = (200, 16, 46)
FONT_DIR = "/usr/share/fonts/truetype/dejavu"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    sizes = {1: 16, 2: 14, 3: 12}
    set_run_font(run, size=sizes.get(level, 11), bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)


def add_body(doc: Document, text: str):
    for para in text.strip().split("\n\n"):
        p = doc.add_paragraph(para.strip())
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            set_run_font(run, size=11)


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, size=11)


def add_image(doc: Document, name: str, width_in: float = 6.0, caption: str | None = None):
    path = IMG / name
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cp = doc.add_paragraph(caption)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cp.runs:
                set_run_font(run, size=9, color=(80, 80, 80))


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Cover
    add_image(doc, "cover_banner.png", width_in=6.5)
    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("DIPLOMA FINAL-YEAR INTERNSHIP\nCOMPANY RESEARCH REPORT")
    set_run_font(r, size=14, bold=True, color=NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = title.add_run("TATA MOTORS LIMITED")
    set_run_font(rt, size=16, bold=True, color=ACCENT)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run(
        "A Study of Company Development, Automotive Products,\n"
        "Research and Development, and Manufacturing Methodologies"
    )
    set_run_font(rs, size=11)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lines = [
        f"Submitted by: {STUDENT['name']}",
        f"Diploma / Department: {STUDENT['department']}",
        f"Institution: {STUDENT['institution']}",
        f"Internship organisation: {STUDENT['internship_org']}",
        f"Academic year: {STUDENT['academic_year']}",
        f"Submission date: {STUDENT['submission_date']}",
        "",
        "Submitted in partial fulfilment of the requirements for the Diploma final-year internship.",
    ]
    for line in lines:
        p = doc.add_paragraph(line)
        for run in p.runs:
            set_run_font(run, size=11)

    doc.add_page_break()

    add_heading(doc, "Certificate / Declaration", 1)
    add_body(
        doc,
        f"I hereby declare that this company research report entitled \"Tata Motors Limited — "
        f"Company Research Report\" is prepared by me, {STUDENT['name']}, as part of the "
        f"final-year diploma internship requirement. The information is compiled from publicly "
        f"available sources and general industrial knowledge. I have not copied any material "
        f"without appropriate acknowledgement.\n\n"
        f"Place: Haveri\nDate: {STUDENT['submission_date']}\n\n"
        f"Signature of student: _______________________\n"
        f"Name: {STUDENT['name']}",
    )

    doc.add_page_break()
    add_heading(doc, "Table of Contents", 1)
    toc = [
        "1. Introduction",
        "2. About Tata Motors",
        "3. Company Overview",
        "4. History and Development of Tata Motors",
        "5. Manufacturing Methodology Used by Tata Motors",
        "6. Products Manufactured by Tata Motors",
        "7. Challenges and Future Opportunities",
        "8. Research Methodology",
        "9. Conclusion",
        "10. References",
    ]
    for line in toc:
        p = doc.add_paragraph(line)
        for run in p.runs:
            set_run_font(run, size=11)

    doc.add_page_break()

    add_heading(doc, "1. Introduction", 1)
    add_body(
        doc,
        "The automobile industry is one of the most important sectors of the Indian economy. It "
        "contributes to industrial development, employment generation, technological innovation, "
        "infrastructure development, and international trade. Within this industry, Tata Motors "
        "has established itself as a major automobile manufacturer with a strong presence in "
        "passenger vehicles, commercial vehicles, electric mobility, and premium international brands.\n\n"
        "Tata Motors was established in 1945 as Tata Engineering and Locomotive Company Limited "
        "(TELCO). Initially, the company focused on locomotives and engineering products. Over "
        "time, it expanded into commercial vehicles, passenger cars, sport utility vehicles (SUVs), "
        "buses, and electric vehicles. The company has contributed significantly to the development "
        "of the Indian automobile industry through indigenous vehicle design, expanded manufacturing "
        "capabilities, improved vehicle safety, and investment in sustainable mobility.\n\n"
        "This research report examines Tata Motors' background, organisational development, business "
        "segments, products, manufacturing processes, research and development areas, achievements, "
        "and recent developments. For a diploma student in Electronics and Communication Engineering, "
        "the report also highlights how electronic systems, sensors, control units, and connectivity "
        "are integrated into modern vehicles.",
    )

    add_heading(doc, "2. About Tata Motors", 1)
    add_body(
        doc,
        "Tata Motors is a leading automotive manufacturer associated with the Tata Group, one of "
        "India's largest and most respected business conglomerates. Its operations cover the design, "
        "development, manufacturing, marketing, and servicing of a wide range of vehicles.\n\n"
        "The company serves individual customers, businesses, logistics operators, public transport "
        "organisations, and industrial customers. Its vehicles are used for personal transportation, "
        "goods transportation, passenger transport, construction, mining, defence-related mobility, "
        "and specialised commercial applications.\n\n"
        "Tata Motors invests in research and development, vehicle electrification, connected "
        "technologies, safety engineering, and advanced manufacturing systems. Public company "
        "materials describe a multi-segment automotive business with strong positions in commercial "
        "vehicles and a growing passenger vehicle and electric mobility portfolio, alongside the "
        "international Jaguar Land Rover (JLR) business.",
    )

    table = doc.add_table(rows=9, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Company name", "Tata Motors Limited (group; segment structure evolved in FY25–26)"),
        ("Established", "1945"),
        ("Original name", "Tata Engineering and Locomotive Company (TELCO)"),
        ("Industry", "Automotive manufacturing"),
        ("Headquarters", "Mumbai, Maharashtra, India"),
        ("Parent group", "Tata Group"),
        ("Main products", "Cars, SUVs, trucks, buses, vans, electric vehicles, mobility solutions"),
        ("Major markets", "India and selected international markets"),
        ("Core focus", "Quality, safety, innovation, sustainability, customer satisfaction"),
    ]
    for i, (a, b) in enumerate(rows):
        table.rows[i].cells[0].text = a
        table.rows[i].cells[1].text = b

    add_heading(doc, "3. Company Overview", 1)
    add_body(
        doc,
        "Tata Motors operates in a highly competitive industry where product quality, pricing, "
        "design, safety, technology, fuel efficiency, and after-sales service influence customer "
        "decisions. According to Tata Motors' integrated annual report and investor presentations "
        "for FY2024–25, the group's automotive business is organised around major reportable "
        "segments including Tata Commercial Vehicles (Tata CV), Tata Passenger Vehicles (Tata PV), "
        "electric mobility, and Jaguar Land Rover (JLR).\n\n"
        "Effective 1 October 2025, the company implemented a composite scheme of arrangement "
        "approved by the National Company Law Tribunal to demerge commercial vehicles from passenger "
        "vehicles (including electric passenger mobility), creating focused listed entities for "
        "sharper strategy execution. Students should refer to the latest annual report for the "
        "exact legal entity names after demerger.",
    )
    add_image(doc, "business_segments.png", 5.8, "Figure 1: Illustrative revenue scale by segment (FY25 public disclosures)")
    add_image(doc, "product_segments.png", 5.8, "Figure 2: Product and market segments")
    add_image(doc, "market_highlights.png", 5.8, "Figure 3: Market position highlights (public FY25 disclosures)")

    add_heading(doc, "3.1 Tata Commercial Vehicles (CV)", 2)
    add_bullets(
        doc,
        [
            "Trucks (light, medium, and heavy), tippers, tractors, and specialised haulage vehicles.",
            "Buses and staff carriers for public transport, schools, and institutions.",
            "Small commercial vehicles and pickups such as the Tata Ace and Intra families.",
            "Electric commercial mobility including Ace EV and electric bus solutions for smart cities.",
            "India market leadership in commercial vehicles; focus on uptime, total cost of ownership, and fleet solutions.",
        ],
    )

    add_heading(doc, "3.2 Tata Passenger Vehicles (PV)", 2)
    add_bullets(
        doc,
        [
            "Hatchbacks and compact cars (e.g., Tiago, Altroz).",
            "SUVs and crossovers (e.g., Punch, Nexon, Harrier, Safari).",
            "Multi-powertrain strategy: petrol, diesel, CNG, and electric variants where offered.",
            "Strong SUV portfolio aligned with Indian market preference; safety ratings (NCAP) used as a brand pillar.",
        ],
    )

    add_heading(doc, "3.3 Electric Mobility", 2)
    add_bullets(
        doc,
        [
            "Tata Passenger Electric Mobility leads India's passenger EV market with a majority share in many periods.",
            "Models include Nexon EV, Tiago EV, Punch EV, and upcoming platforms.",
            "Ecosystem work includes charging partnerships, battery technology, and software-defined features.",
            "Commercial EV includes electric buses with large fleet kilometres operated in Indian cities.",
        ],
    )

    add_heading(doc, "3.4 Jaguar Land Rover (JLR)", 2)
    add_body(
        doc,
        "Acquired in 2008, JLR represents Tata Motors' global premium automotive business with "
        "luxury brands Jaguar and Land Rover, advanced engineering, and export-oriented manufacturing "
        "in the United Kingdom and other locations.",
    )

    add_heading(doc, "4. History and Development of Tata Motors", 1)
    add_image(doc, "history_timeline.png", 6.2, "Figure 4: Key milestones in Tata Motors history")
    add_body(
        doc,
        "1945 — Tata Engineering and Locomotive Company (TELCO) founded; early focus on locomotives "
        "and engineering.\n\n"
        "1950s — Entry into commercial vehicles through collaboration with Daimler-Benz.\n\n"
        "1998 — Launch of Tata Indica, an important indigenous passenger car milestone.\n\n"
        "2004 — Acquisition of Daewoo Commercial Vehicle Company (South Korea), expanding international CV footprint.\n\n"
        "2008 — Acquisition of Jaguar Land Rover, transforming global brand and technology access.\n\n"
        "2010s–2020s — Modern PV lineup (Tiago, Nexon, Harrier, Punch), safety focus, and rapid EV scaling.\n\n"
        "2025 — Business demerger separating CV and PV strategies; continued investment in EV, CNG, and digital manufacturing.",
    )

    add_heading(doc, "5. Manufacturing Methodology Used by Tata Motors", 1)
    add_image(doc, "manufacturing_flow.png", 6.2, "Figure 5: Typical vehicle manufacturing flow")
    add_body(
        doc,
        "Automobile manufacturing involves coordinated activities from product design and procurement "
        "to assembly, inspection, and delivery. Tata Motors uses established automotive practices; "
        "exact procedures vary by plant and product line. The following describes widely used "
        "methodologies relevant to Tata Motors and the industry.",
    )
    sections_mfg = [
        (
            "5.1 Lean manufacturing",
            "Reduces waste and improves material flow through elimination of unnecessary movement, "
            "excess inventory, and rework while improving productivity and quality.",
        ),
        (
            "5.2 Just-in-Time (JIT) production",
            "Supplies components when needed on the line, requiring reliable suppliers and accurate scheduling.",
        ),
        (
            "5.3 Kaizen and continuous improvement",
            "Employees and supervisors identify problems and implement incremental improvements in safety, quality, and workflow.",
        ),
        (
            "5.4 Total Quality Management (TQM)",
            "Quality is built across design, production, and service—not only at final inspection.",
        ),
        (
            "5.5 Automation, robotics, CAD/CAM",
            "Robots improve precision in welding and handling; CAD/CAM supports digital design and manufacturing planning.",
        ),
        (
            "5.6 Supply chain and vendor management",
            "Steel, tyres, electronics, batteries, and assemblies are sourced from a large vendor network with incoming quality checks.",
        ),
        (
            "5.7 Safety and environmental management",
            "Plants manage occupational safety, energy, emissions, wastewater, and statutory environmental compliance.",
        ),
        (
            "5.8 Electric vehicle manufacturing",
            "Adds high-voltage battery assembly, thermal management, power electronics, and strict electrical safety protocols.",
        ),
        (
            "5.9 Quality inspection and testing",
            "Includes body and paint checks, electrical tests, powertrain checks, brake/steering inspection, leak tests, and road tests as applicable.",
        ),
    ]
    for title, text in sections_mfg:
        add_heading(doc, title, 2)
        add_body(doc, text)

    add_heading(doc, "6. Products Manufactured by Tata Motors", 1)
    add_body(
        doc,
        "Tata Motors manufactures a wide portfolio for personal mobility, logistics, and public transport. "
        "Representative families are summarised below (model lists change with product lifecycle).",
    )
    prod_table = doc.add_table(rows=6, cols=3)
    prod_table.style = "Table Grid"
    hdr = prod_table.rows[0].cells
    hdr[0].text = "Segment"
    hdr[1].text = "Category"
    hdr[2].text = "Examples"
    data = [
        ("Passenger", "Hatch / sedan", "Tiago, Altroz, Tigor"),
        ("Passenger", "SUV / crossover", "Punch, Nexon, Harrier, Safari, Curvv"),
        ("Commercial", "SCV / LCV / M&HCV", "Ace, Intra, Ultra, Prima, Signa"),
        ("Commercial", "Buses & mobility", "Starbus, electric buses, staff carriers"),
        ("Electric", "Passenger & CV EV", "Nexon EV, Tiago EV, Punch EV, Ace EV"),
    ]
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            prod_table.rows[i].cells[j].text = val

    add_image(doc, "vehicle_portfolio.png", 6.2, "Figure 6: Representative product families across segments")
    add_image(doc, "ev_ecosystem.png", 5.8, "Figure 7: Electric mobility value chain")

    add_heading(doc, "6.1 Research, Development and Electronics (ECE relevance)", 2)
    add_body(
        doc,
        "Modern automobiles are complex electro-mechanical systems. Electronics and Communication "
        "Engineering plays a critical role in automotive R&D and production quality.\n\n"
        "Key electronic domains in Tata Motors vehicles include engine control units (ECU), body "
        "control modules, infotainment and telematics, advanced driver assistance sensors, battery "
        "management systems (BMS) in EVs, wiring harnesses, communication buses (CAN, LIN, Ethernet "
        "in newer architectures), and over-the-air (OTA) software updates in connected cars.\n\n"
        "Manufacturing and quality teams use electronic test equipment for functional verification, "
        "diagnostics, and traceability. Diploma interns in electronics may observe PLC-based automation, "
        "robotic cells, end-of-line electrical checks, and instrumentation in paint and assembly shops "
        "during authorised plant visits.",
    )
    add_bullets(
        doc,
        [
            "Embedded systems: microcontrollers in ECUs and BMS.",
            "Sensors: speed, temperature, pressure, radar/camera for ADAS on higher trims.",
            "Communication: in-vehicle networks linking sensors, displays, and control units.",
            "Power electronics: inverters and converters in electric powertrains.",
            "Quality electronics: test jigs, oscilloscopes, and automated diagnostic stations.",
        ],
    )
    add_image(doc, "ece_automotive.png", 5.5, "Figure 8: Electronics integration in modern vehicles (ECE perspective)")

    add_heading(doc, "6.2 Manufacturing Plants and Operations in India", 2)
    add_body(
        doc,
        "Tata Motors operates multiple manufacturing and assembly facilities in India. Major locations "
        "include Pune (passenger vehicles), Jamshedpur and Lucknow (commercial vehicles), Sanand "
        "(passenger and EV-related capacity), and Dharwad among others. Plants integrate stamping, "
        "body shop, paint shop, engine or powertrain assembly (where applicable), final assembly, "
        "and quality gates. Capacity and product mix are updated over time as new models launch.",
    )
    add_image(doc, "india_plants_map.png", 5.2, "Figure 9: Major Tata Motors locations in India (schematic)")

    add_heading(doc, "7. Challenges and Future Opportunities", 1)
    add_bullets(
        doc,
        [
            "Competition and pricing pressure across CV and PV segments.",
            "Supply chain resilience for semiconductors, batteries, and specialised components.",
            "Quality and reliability expectations from customers and regulators.",
            "Accelerating electrification, charging infrastructure, and battery cost reduction.",
            "Sustainability, carbon footprint, and compliance with emission and safety norms.",
            "Opportunities in connected vehicles, software-defined features, hydrogen/fuel-cell research, and export growth.",
        ],
    )

    add_heading(doc, "8. Research Methodology", 1)
    add_body(
        doc,
        "This report uses a descriptive, document-based research approach. Publicly available company "
        "information and general automotive manufacturing knowledge are organised to explain Tata Motors' "
        "development, segments, products, R&D areas, manufacturing methodologies, and challenges.",
    )
    add_bullets(
        doc,
        [
            "Secondary data: Tata Motors website, integrated annual report FY2024–25, investor presentations, product catalogues, and reputable industry publications.",
            "Primary data (if collected during internship): authorised plant visit notes, supervised discussions, and observation logs—include only with official permission.",
            "Analysis: historical timeline, segment-wise description, and thematic grouping (quality, manufacturing, sustainability, electronics).",
            "Verification: cross-check figures with source publication dates and financial year labels.",
            "Limitations: confidential plant procedures and unreleased R&D details are not available in public sources.",
        ],
    )

    add_heading(doc, "9. Conclusion", 1)
    add_body(
        doc,
        "Tata Motors has grown from a locomotive and engineering company into a diversified automotive "
        "group with leadership in Indian commercial vehicles, a strong passenger vehicle portfolio, "
        "market-leading electric mobility initiatives, and a global premium business through Jaguar Land Rover. "
        "Its manufacturing philosophy combines lean production, quality systems, automation, and supply chain "
        "coordination. For diploma students, studying Tata Motors provides practical insight into how "
        "large-scale manufacturing, electronics integration, and innovation align with national industrial "
        "growth and sustainable mobility goals.",
    )

    add_heading(doc, "10. References", 1)
    refs = [
        "Tata Motors Limited, Integrated Annual Report 2024–25.",
        "Tata Motors Limited, Q4 FY25 Investor Presentation and press releases.",
        "Tata Motors Limited, Corporate website — www.tatamotors.com (company profile, products, sustainability).",
        "Government of India, automotive industry reports and policy documents (for sector context).",
        "Textbooks on automobile engineering and manufacturing technology (for general methodology terms).",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph(f"[{i}] {ref}")
        for run in p.runs:
            set_run_font(run, size=10)

    out_path = OUT / "TATA_Motors_Research_Report_Nachiket_L_Karoshi.docx"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


class ReportPDF(FPDF):
    def header(self):
        self.set_font("DejaVu", "B", 9)
        self.set_text_color(0, 51, 102)
        self.cell(0, 8, "Tata Motors | Company Research Report", align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(1)

    def footer(self):
        self.set_y(-12)
        self.set_font("DejaVu", "", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")


def pdf_add_section(pdf: ReportPDF, title: str, body: str, img: str | None = None):
    pdf.set_x(pdf.l_margin)
    pdf.set_font("DejaVu", "B", 12)
    pdf.set_text_color(0, 51, 102)
    pdf.multi_cell(pdf.epw, 7, title)
    pdf.ln(2)
    pdf.set_x(pdf.l_margin)
    pdf.set_font("DejaVu", "", 10)
    pdf.set_text_color(30, 30, 30)
    pdf.multi_cell(pdf.epw, 5.5, body)
    pdf.ln(3)
    if img and (IMG / img).exists():
        w = min(170, pdf.epw)
        if pdf.get_y() > 230:
            pdf.add_page()
        x = pdf.l_margin + (pdf.epw - w) / 2
        pdf.image(str(IMG / img), x=x, w=w)
        pdf.ln(6)
        pdf.set_x(pdf.l_margin)


def build_pdf() -> Path:
    pdf = ReportPDF()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_font("DejaVu", "", f"{FONT_DIR}/DejaVuSans.ttf")
    pdf.add_font("DejaVu", "B", f"{FONT_DIR}/DejaVuSans-Bold.ttf")
    pdf.add_page()

    if (IMG / "cover_banner.png").exists():
        pdf.image(str(IMG / "cover_banner.png"), x=10, y=18, w=190)
        pdf.set_y(95)

    pdf.set_x(pdf.l_margin)
    pdf.set_font("DejaVu", "B", 14)
    pdf.multi_cell(pdf.epw, 8, "DIPLOMA FINAL-YEAR INTERNSHIP\nCOMPANY RESEARCH REPORT", align="C")
    pdf.set_font("DejaVu", "B", 13)
    pdf.multi_cell(pdf.epw, 7, "TATA MOTORS LIMITED", align="C")
    pdf.set_font("DejaVu", "", 10)
    pdf.multi_cell(
        pdf.epw,
        5.5,
        "A Study of Company Development, Automotive Products,\n"
        "Research and Development, and Manufacturing Methodologies",
        align="C",
    )
    pdf.ln(3)
    cover = (
        f"Submitted by: {STUDENT['name']}\n"
        f"Diploma / Department: {STUDENT['department']}\n"
        f"Institution: {STUDENT['institution']}\n"
        f"Internship organisation: {STUDENT['internship_org']}\n"
        f"Academic year: {STUDENT['academic_year']}\n"
        f"Submission date: {STUDENT['submission_date']}\n\n"
        "Submitted in partial fulfilment of the requirements for the Diploma final-year internship."
    )
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(pdf.epw, 6, cover)
    pdf.add_page()

    sections = [
        (
            "1. Introduction",
            "The automobile industry is one of the most important sectors of the Indian economy. It "
            "contributes to industrial development, employment, innovation, infrastructure, and trade. "
            "Tata Motors is a major manufacturer in passenger vehicles, commercial vehicles, electric "
            "mobility, and engineering. Established in 1945 as TELCO, it expanded from locomotives to "
            "cars, SUVs, trucks, buses, and EVs. This report studies background, segments, products, "
            "manufacturing, R&D, and challenges.",
            None,
        ),
        (
            "2. About Tata Motors",
            "Tata Motors is part of the Tata Group. It designs, develops, manufactures, markets, and "
            "services vehicles for individuals, fleets, logistics, and public transport. It invests in "
            "R&D, electrification, connected technology, safety, and advanced manufacturing.",
            None,
        ),
        (
            "3. Company Overview",
            "Business segments (FY25): Tata Commercial Vehicles, Tata Passenger Vehicles, electric "
            "mobility, and Jaguar Land Rover. From October 2025, CV and PV were demerged into focused "
            "listed companies (NCLT-approved). CV includes trucks, buses, Ace/Intra; PV includes "
            "Tiago, Nexon, Harrier; EV includes Nexon EV and e-buses; JLR is the global luxury business.",
            "business_segments.png",
        ),
        ("", "", "product_segments.png"),
        ("", "", "market_highlights.png"),
        (
            "4. History and Development of Tata Motors",
            "1945: TELCO founded. 1950s: commercial vehicles with Daimler-Benz. 1998: Tata Indica. "
            "2004: Daewoo CV acquisition. 2008: Jaguar Land Rover acquisition. 2010s-2020s: Tiago, "
            "Nexon, Harrier, Punch, EV leadership. 2025: CV-PV demerger and digital manufacturing.",
            "history_timeline.png",
        ),
        (
            "5. Manufacturing Methodology Used by Tata Motors",
            "Lean manufacturing, Just-in-Time, Kaizen, TQM, automation and robotics, CAD/CAM, supply "
            "chain management, environmental and safety management, EV high-voltage processes, and "
            "end-of-line quality testing including electrical and road tests.",
            "manufacturing_flow.png",
        ),
        (
            "6. Products Manufactured by Tata Motors",
            "Passenger: Tiago, Altroz, Tigor, Punch, Nexon, Harrier, Safari, Curvv. Commercial: Ace, "
            "Intra, Yodha, Ultra, Signa, Prima, Starbus. Electric: Nexon EV, Tiago EV, Punch EV, Ace EV, "
            "electric buses. Products serve personal mobility, goods transport, and public transport.",
            "vehicle_portfolio.png",
        ),
        ("", "", "ev_ecosystem.png"),
        (
            "6.1 R&D and Electronics (ECE)",
            "Modern vehicles use ECUs, BCM, BMS (EV), sensors, infotainment, telematics, CAN/LIN buses, "
            "and OTA updates. ECE students study embedded systems, power electronics, and production testing.",
            "ece_automotive.png",
        ),
        (
            "6.2 Plants in India",
            "Major plants: Pune (PV), Jamshedpur and Lucknow (CV), Sanand (PV/EV), Dharwad (CV). "
            "Processes include stamping, body shop, paint, assembly, and quality gates.",
            "india_plants_map.png",
        ),
        (
            "7. Challenges and Future Opportunities",
            "Challenges: competition, pricing, supply chain, quality, electrification, regulations. "
            "Opportunities: R&D, digitalisation, energy efficiency, battery recycling, connected vehicles.",
            None,
        ),
        (
            "8. Research Methodology",
            "Descriptive document-based research using annual reports, investor presentations, product "
            "catalogues, and authorised internship observations. Limitations: confidential plant data "
            "is not publicly available.",
            None,
        ),
        (
            "9. Conclusion",
            "Tata Motors is a diversified automotive leader in India with strong CV and PV portfolios, "
            "EV leadership, and global presence through JLR. Its manufacturing and electronics integration "
            "offer valuable learning for diploma internship and industrial visits.",
            None,
        ),
    ]
    for title, body, image in sections:
        if title or body:
            pdf_add_section(pdf, title, body, None)
        if image:
            w = min(170, pdf.epw)
            if pdf.get_y() > 230:
                pdf.add_page()
            x = pdf.l_margin + (pdf.epw - w) / 2
            pdf.image(str(IMG / image), x=x, w=w)
            pdf.ln(6)
            pdf.set_x(pdf.l_margin)

    pdf_add_section(
        pdf,
        "10. References",
        "[1] Tata Motors Integrated Annual Report 2024-25.\n"
        "[2] Tata Motors Q4 FY25 investor presentation.\n"
        "[3] www.tatamotors.com\n"
        "[4] Government of India automotive sector publications.",
        None,
    )

    out_path = OUT / "TATA_Motors_Research_Report_Nachiket_L_Karoshi.pdf"
    pdf.output(str(out_path))
    return out_path


def package_downloads(docx: Path, pdf: Path) -> Path:
    import shutil
    import zipfile

    aliases = [
        ("TATA_Motors_Diploma_Project.docx", docx),
        ("TATA_Motors_Diploma_Project.pdf", pdf),
    ]
    for name, src in aliases:
        shutil.copy2(src, OUT / name)

    zip_path = OUT / "TATA_Motors_Diploma_Project_DOWNLOAD.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, src in aliases:
            zf.write(src, arcname=name)
    return zip_path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    docx = build_docx()
    pdf = build_pdf()
    zip_path = package_downloads(docx, pdf)
    print(f"DOCX: {docx}")
    print(f"PDF:  {pdf}")
    print(f"ZIP:  {zip_path}")


if __name__ == "__main__":
    main()
